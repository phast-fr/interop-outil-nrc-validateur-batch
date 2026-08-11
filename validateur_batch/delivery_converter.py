"""Reconditionnement des classeurs de livraison selon la règle REMP ASCT-188.

Le classeur original n'est jamais modifié. Une copie reconditionnée et les CSV
compacts attendus par :class:`validateur_batch.object.batch.Batch` sont écrits
dans le dossier de sortie.
"""

from __future__ import annotations

import csv
import re
import tempfile
import warnings
import zipfile
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from validateur_batch.object.batch import COL


SHEET_BY_TYPE = {
    "ADD": "Description Additions",
    "CHG": "Description Changes",
    "INA": "Description Inactivations",
    "REP": "Description Replacements",
}

KEY_COLUMN = {
    "ADD": "Concept ID",
    "CHG": "Description ID",
    "INA": "Description ID Or Term",
    "REP": "Concept ID",
}

OUTPUT_NAME = {
    "ADD": "Descriptions_Additions.csv",
    "CHG": "Descriptions_Changes.csv",
    "INA": "Descriptions_Inactivations.csv",
    "REP": "Descriptions_Replacements.csv",
}

REPORT_COLUMNS = [
    "file",
    "sheet",
    "row",
    "conceptId",
    "oldDescriptionId",
    "newDescriptionId",
    "action",
    "status",
    "message",
]


class DeliveryConversionError(ValueError):
    """Erreur bloquante détectée pendant le reconditionnement."""


@dataclass(frozen=True)
class PreparedDelivery:
    """Fichiers produits par le reconditionnement d'une livraison."""

    batch_files: dict[str, str]
    workbook: str
    report: str
    error_report: str | None
    rejected_rows: int


def _text(value: object) -> str:
    """Retourne une valeur Excel normalisée en texte."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _headers(ws: Worksheet) -> list[str]:
    return [_text(cell.value) for cell in ws[1]]


def _rows(ws: Worksheet, key_header: str) -> list[dict[str, object]]:
    """Lit uniquement les lignes métier, sans les lignes préformatées vides."""
    headers = _headers(ws)
    try:
        key_index = headers.index(key_header)
    except ValueError as exc:
        raise DeliveryConversionError(
            f"Colonne obligatoire absente de l'onglet {ws.title!r}: {key_header}"
        ) from exc

    result = []
    for excel_row in range(2, ws.max_row + 1):
        values = [ws.cell(excel_row, col + 1).value for col in range(len(headers))]
        if not _text(values[key_index]):
            continue
        result.append(
            {
                "excel_row": excel_row,
                "values": values,
                "data": dict(zip(headers, values)),
            }
        )
    return result


def _description(
    descriptions: pd.DataFrame,
    description_id: str,
    role: str,
) -> pd.Series:
    matches = descriptions.loc[descriptions["id"].astype(str) == description_id]
    if matches.empty:
        raise DeliveryConversionError(
            f"La {role} {description_id} est absente des descriptions "
            "françaises actives."
        )

    values = matches[
        ["conceptId", "term", "caseSignificanceId", "acceptabilityId"]
    ].drop_duplicates()
    if len(values) != 1:
        raise DeliveryConversionError(
            f"La {role} {description_id} possède plusieurs états actifs incohérents."
        )
    return values.iloc[0]


def _row_from_headers(headers: list[str], values: dict[str, object]) -> list[object]:
    """Construit une ligne complète, y compris les paires de refsets répétées."""
    used: dict[str, int] = {}
    row = []
    for header in headers:
        occurrence = used.get(header, 0)
        candidate = values.get(header, "") if occurrence == 0 else ""
        row.append(candidate)
        used[header] = occurrence + 1
    return row


def _copy_row_style(ws: Worksheet, source_row: int, target_row: int) -> None:
    """Prolonge le style si le classeur n'est pas préformaté."""
    if source_row == target_row:
        return
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style and not target.has_style:
            target._style = copy(source._style)
        if source.number_format and not target.number_format:
            target.number_format = source.number_format


def _write_business_rows(ws: Worksheet, rows: Iterable[list[object]]) -> None:
    """Réécrit les lignes métier sans supprimer la mise en forme du template."""
    rows = list(rows)
    populated = _rows(ws, _headers(ws)[0])
    last_populated = max((item["excel_row"] for item in populated), default=1)
    last_target = 1 + len(rows)
    last_clear = max(last_populated, last_target)

    for row_index in range(2, last_clear + 1):
        if row_index > ws.max_row:
            _copy_row_style(ws, row_index - 1, row_index)
        values = rows[row_index - 2] if row_index <= last_target else []
        for col in range(1, ws.max_column + 1):
            value = values[col - 1] if col <= len(values) else None
            ws.cell(row_index, col).value = value


def _compact_dataframe(ws: Worksheet, batch_type: str) -> pd.DataFrame:
    """Extrait les premières colonnes de langue attendues par Batch."""
    headers = _headers(ws)
    indices = []
    start = 0
    for expected in COL[batch_type]:
        try:
            index = headers.index(expected, start)
        except ValueError as exc:
            raise DeliveryConversionError(
                f"Colonne {expected!r} absente de l'onglet {ws.title!r}."
            ) from exc
        indices.append(index)
        start = index + 1

    records = []
    for item in _rows(ws, KEY_COLUMN[batch_type]):
        records.append([_text(item["values"][index]) for index in indices])
    return pd.DataFrame(records, columns=COL[batch_type], dtype=str)


def _worksheet_paths(archive: zipfile.ZipFile) -> dict[str, str]:
    """Associe chaque nom d'onglet à son fichier XML dans le classeur."""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document_rel_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationships = ElementTree.fromstring(
        archive.read("xl/_rels/workbook.xml.rels")
    )
    targets = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in relationships.findall(f"{{{package_rel_ns}}}Relationship")
    }
    result = {}
    for sheet in workbook.findall(f"{{{main_ns}}}sheets/{{{main_ns}}}sheet"):
        target = targets[sheet.attrib[f"{{{document_rel_ns}}}id"]].replace("\\", "/")
        if target.startswith("/"):
            target = target.removeprefix("/")
        elif not target.startswith("xl/"):
            target = f"xl/{target}"
        result[sheet.attrib["name"]] = target
    return result


def _replace_sheet_data(original: bytes, modified: bytes) -> bytes:
    """Remplace uniquement ``sheetData`` en conservant tout le reste du template."""
    pattern = re.compile(
        rb"<(?:[A-Za-z0-9_]+:)?sheetData\b.*?"
        rb"</(?:[A-Za-z0-9_]+:)?sheetData>",
        re.DOTALL,
    )
    replacement = pattern.search(modified)
    if replacement is None or pattern.search(original) is None:
        raise DeliveryConversionError("Structure OOXML sheetData introuvable.")
    return pattern.sub(replacement.group(0), original, count=1)


def _save_preserving_template(
    source: Path,
    modified: Path,
    target: Path,
    modified_sheets: Iterable[str],
) -> None:
    """Crée le livrable depuis l'archive originale sans dégrader le template."""
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(modified) as modified_zip:
        source_paths = _worksheet_paths(source_zip)
        modified_paths = _worksheet_paths(modified_zip)
        replacements = {
            source_paths[name]: modified_zip.read(modified_paths[name])
            for name in modified_sheets
        }
        with tempfile.NamedTemporaryFile(
            dir=target.parent, suffix=".xlsx", delete=False
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        try:
            with zipfile.ZipFile(temporary_path, "w") as rebuilt:
                for item in source_zip.infolist():
                    content = source_zip.read(item.filename)
                    if item.filename in replacements:
                        content = _replace_sheet_data(
                            content, replacements[item.filename]
                        )
                    rebuilt.writestr(item, content)
            temporary_path.replace(target)
        finally:
            temporary_path.unlink(missing_ok=True)


def prepare_delivery_inputs(
    delivery_file: str,
    descriptions: pd.DataFrame,
    output_dir: str,
) -> PreparedDelivery:
    """Reconditionne un classeur puis produit les entrées du validateur.

    Les REMP avec ``New Translated Term`` restent dans REP. Les REMP pointant
    vers une description existante deviennent une promotion CHG et une
    inactivation INA. Toutes les lignes sont analysées avant d'écrire les
    résultats définitifs.
    """
    source = Path(delivery_file)
    if not source.is_file():
        raise DeliveryConversionError(f"Fichier de livraison introuvable: {source}")

    destination = Path(output_dir) / "recreated_inputs"
    destination.mkdir(parents=True, exist_ok=True)
    report_path = destination / "conversion_report.csv"
    error_report_path = destination / "conversion_errors.csv"

    try:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore", message="Data Validation extension is not supported"
            )
            workbook = load_workbook(source)
    except Exception as exc:
        raise DeliveryConversionError(
            f"Impossible de lire le fichier de livraison {source}: {exc}"
        ) from exc

    missing = [
        name for name in SHEET_BY_TYPE.values() if name not in workbook.sheetnames
    ]
    if missing:
        raise DeliveryConversionError(
            "Onglet(s) obligatoire(s) absent(s): " + ", ".join(missing)
        )

    sheets = {kind: workbook[name] for kind, name in SHEET_BY_TYPE.items()}
    rows_by_type = {
        kind: _rows(sheet, KEY_COLUMN[kind]) for kind, sheet in sheets.items()
    }
    report: list[dict[str, object]] = []
    generated_chg: list[list[object]] = []
    generated_ina: list[list[object]] = []
    kept_rep: list[list[object]] = []

    chg_ids = {
        _text(item["data"].get("Description ID")) for item in rows_by_type["CHG"]
    }
    ina_ids = {
        _text(item["data"].get("Description ID Or Term"))
        for item in rows_by_type["INA"]
    }

    for item in rows_by_type["REP"]:
        data = item["data"]
        concept_id = _text(data.get("Concept ID"))
        old_id = _text(data.get("Description ID"))
        new_id = _text(data.get("New Replacement Description ID"))
        new_term = _text(data.get("New Translated Term"))
        base_report = {
            "file": source.name,
            "sheet": SHEET_BY_TYPE["REP"],
            "row": item["excel_row"],
            "conceptId": concept_id,
            "oldDescriptionId": old_id,
            "newDescriptionId": new_id,
        }

        if bool(new_id) == bool(new_term):
            message = (
                "Renseigner soit New Replacement Description ID, soit "
                "New Translated Term, mais pas les deux."
            )
            report.append(
                {**base_report, "action": "NONE", "status": "ERROR", "message": message}
            )
            continue

        if new_term:
            kept_rep.append(item["values"])
            report.append(
                {
                    **base_report,
                    "action": "REP_KEPT",
                    "status": "OK",
                    "message": (
                        "Remplacement par une nouvelle description conservé dans REP."
                    ),
                }
            )
            continue

        try:
            if old_id == new_id:
                raise DeliveryConversionError(
                    "L'ancien et le nouveau Description ID sont identiques."
                )
            old_desc = _description(descriptions, old_id, "description à inactiver")
            new_desc = _description(descriptions, new_id, "description à promouvoir")
            if _text(old_desc["conceptId"]) != concept_id:
                raise DeliveryConversionError(
                    "L'ancienne description appartient au concept "
                    f"{old_desc['conceptId']} "
                    f"et non au concept {concept_id}."
                )
            if _text(new_desc["conceptId"]) != concept_id:
                raise DeliveryConversionError(
                    "La nouvelle description appartient au concept "
                    f"{new_desc['conceptId']} "
                    f"et non au concept {concept_id}."
                )
            if _text(old_desc["acceptabilityId"]) != "PREFERRED":
                raise DeliveryConversionError(
                    "La description à inactiver n'est pas le PT français actif."
                )
            if _text(new_desc["acceptabilityId"]) != "ACCEPTABLE":
                raise DeliveryConversionError(
                    "La description à promouvoir n'est pas un synonyme "
                    "acceptable actif."
                )
            if new_id in chg_ids:
                raise DeliveryConversionError(
                    f"La description {new_id} est déjà présente dans "
                    "Description Changes."
                )

            chg_values = {
                "Description ID": new_id,
                "Preferred Term (For reference only)": data.get(
                    "Preferred Term (For reference only)", ""
                ),
                "Term (For reference only)": _text(new_desc["term"]),
                "Case significance": _text(new_desc["caseSignificanceId"]),
                "Type": "SYNONYM",
                "Language reference set": "French",
                "Acceptability": "PREFERRED",
                "Notes": data.get("Notes", ""),
            }
            if old_id not in ina_ids:
                ina_values = {
                    "Description ID Or Term": old_id,
                    "Language Code (require if the term is specified)": "fr",
                    "Concept ID (Optional)": concept_id,
                    "Preferred Term (For reference only)": data.get(
                        "Preferred Term (For reference only)", ""
                    ),
                    "Term (For reference only)": _text(old_desc["term"]),
                    "Inactivation Reason": data.get("Inactivation Reason", ""),
                    "Association Target ID1": data.get("Association Target ID1", ""),
                    "Association Target ID2": data.get("Association Target ID2", ""),
                    "Association Target ID3": data.get("Association Target ID3", ""),
                    "Association Target ID4": data.get("Association Target ID4", ""),
                    "Notes": data.get("Notes", ""),
                }
                generated_ina.append(
                    _row_from_headers(_headers(sheets["INA"]), ina_values)
                )
                ina_ids.add(old_id)
            generated_chg.append(_row_from_headers(_headers(sheets["CHG"]), chg_values))
            chg_ids.add(new_id)
            report.append(
                {
                    **base_report,
                    "action": "REP_TO_CHG_INA",
                    "status": "OK",
                    "message": (
                        "Description existante promue dans CHG et ancien PT "
                        "inactivé dans INA."
                    ),
                }
            )
        except DeliveryConversionError as exc:
            report.append(
                {
                    **base_report,
                    "action": "NONE",
                    "status": "ERROR",
                    "message": str(exc),
                }
            )

    report_dataframe = pd.DataFrame(report, columns=REPORT_COLUMNS)
    report_dataframe.to_csv(
        report_path, sep=";", index=False, encoding="utf-8-sig"
    )
    error_dataframe = report_dataframe.loc[report_dataframe["status"] == "ERROR"]
    if not error_dataframe.empty:
        error_dataframe.to_csv(
            error_report_path, sep=";", index=False, encoding="utf-8-sig"
        )
    else:
        error_report_path.unlink(missing_ok=True)

    original_chg = [item["values"] for item in rows_by_type["CHG"]]
    original_ina = [item["values"] for item in rows_by_type["INA"]]
    _write_business_rows(sheets["CHG"], original_chg + generated_chg)
    _write_business_rows(sheets["INA"], original_ina + generated_ina)
    _write_business_rows(sheets["REP"], kept_rep)

    workbook_path = destination / f"{source.stem}_reconditionne.xlsx"
    with tempfile.NamedTemporaryFile(
        dir=destination, suffix=".xlsx", delete=False
    ) as modified_file:
        modified_path = Path(modified_file.name)
    try:
        workbook.save(modified_path)
        _save_preserving_template(
            source,
            modified_path,
            workbook_path,
            (SHEET_BY_TYPE["CHG"], SHEET_BY_TYPE["INA"], SHEET_BY_TYPE["REP"]),
        )
    finally:
        modified_path.unlink(missing_ok=True)

    batch_files: dict[str, str] = {}
    for kind, sheet in sheets.items():
        dataframe = _compact_dataframe(sheet, kind)
        path = destination / OUTPUT_NAME[kind]
        dataframe.to_csv(
            path,
            sep=";",
            index=False,
            encoding="utf-8-sig",
            quoting=csv.QUOTE_NONE,
            escapechar="\\",
        )
        if not dataframe.empty:
            batch_files[kind] = str(path)

    return PreparedDelivery(
        batch_files=batch_files,
        workbook=str(workbook_path),
        report=str(report_path),
        error_report=(str(error_report_path) if not error_dataframe.empty else None),
        rejected_rows=len(error_dataframe),
    )
