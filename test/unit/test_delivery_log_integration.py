from pathlib import Path

from openpyxl import Workbook, load_workbook

from validateur_batch.delivery_log_integration import append_delivery_errors_to_log


def test_append_delivery_errors_to_log(tmp_path: Path) -> None:
    logs_path = tmp_path / "Logs 2026-08-06 04-17.xlsx"
    workbook = Workbook()
    logs_sheet = workbook.active
    logs_sheet.title = "Logs"
    logs_sheet.append(["ConceptId", "FsnEn", "Message", "Type", "N° ligne", "Fichier"])
    logs_sheet.append(["10", "Kidney structure", "erreur outil", "ERROR", "3", "in.xlsx"])
    workbook.save(logs_path)

    errors_csv = tmp_path / "conversion_errors.csv"
    errors_csv.write_text(
        "file;sheet;row;conceptId;oldDescriptionId;newDescriptionId;action;status;message\n"
        "livraison.xlsx;Description Replacements;5;10;100;200;NONE;ERROR;desc absente\n",
        encoding="utf-8-sig",
    )

    append_delivery_errors_to_log(str(errors_csv), str(logs_path))

    result = load_workbook(logs_path)
    assert "Logs" in result.sheetnames
    assert "Logs REMP" in result.sheetnames

    remp_rows = list(result["Logs REMP"].iter_rows(values_only=True))
    assert remp_rows[0] == ("ConceptId", "FsnEn", "Message", "Type", "N° ligne", "Fichier")
    assert remp_rows[1] == ("10", None, "desc absente", "ERROR", "5", "livraison.xlsx")

    # L'onglet original de l'outil de transformation reste intact.
    original_rows = list(result["Logs"].iter_rows(values_only=True))
    assert original_rows[1] == ("10", "Kidney structure", "erreur outil", "ERROR", "3", "in.xlsx")


def test_append_delivery_errors_to_log_replaces_existing_sheet(tmp_path: Path) -> None:
    logs_path = tmp_path / "Logs.xlsx"
    workbook = Workbook()
    workbook.active.title = "Logs"
    workbook.create_sheet("Logs REMP").append(["stale"])
    workbook.save(logs_path)

    errors_csv = tmp_path / "conversion_errors.csv"
    errors_csv.write_text(
        "file;sheet;row;conceptId;oldDescriptionId;newDescriptionId;action;status;message\n"
        "livraison.xlsx;Description Replacements;5;10;100;200;NONE;ERROR;desc absente\n",
        encoding="utf-8-sig",
    )

    append_delivery_errors_to_log(str(errors_csv), str(logs_path))

    result = load_workbook(logs_path)
    remp_rows = list(result["Logs REMP"].iter_rows(values_only=True))
    assert remp_rows[0] == ("ConceptId", "FsnEn", "Message", "Type", "N° ligne", "Fichier")
    assert len(remp_rows) == 2
