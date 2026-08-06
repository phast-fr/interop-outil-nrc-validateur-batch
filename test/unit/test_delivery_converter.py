from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from validateur_batch.delivery_converter import (
    DeliveryConversionError,
    prepare_delivery_inputs,
)
from validateur_batch.object.batch import COL


SHEETS = {
    "ADD": "Description Additions",
    "CHG": "Description Changes",
    "INA": "Description Inactivations",
    "REP": "Description Replacements",
}


def _delivery(path: Path, replacement_rows: list[list[str]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for batch_type, sheet_name in SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(COL[batch_type] + (["Notes"] if batch_type != "ADD" else []))
        if batch_type == "REP":
            for row in replacement_rows:
                sheet.append(row)
    workbook.save(path)


def _rep_row(new_id: str = "200", new_term: str = "") -> list[str]:
    values = {
        "Concept ID": "10",
        "Description ID": "100",
        "Preferred Term (For reference only)": "English PT",
        "Term (For reference only)": "ancien PT",
        "Inactivation Reason": "Non-conformance to editorial policy",
        "New Replacement Description ID": new_id,
        "Replacement term (For reference only)": "synonyme existant",
        "New Translated Term": new_term,
        "Language Code": "fr" if new_term else "",
        "Case significance": "ci" if new_term else "",
        "Type": "SYNONYM" if new_term else "",
        "Language reference set": "French" if new_term else "",
        "Acceptability": "PREFERRED" if new_term else "",
        "Notes": "note",
    }
    return [values.get(header, "") for header in COL["REP"] + ["Notes"]]


@pytest.fixture
def descriptions() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "id": "100",
                "active": "1",
                "conceptId": "10",
                "term": "ancien PT",
                "caseSignificanceId": "ci",
                "acceptabilityId": "PREFERRED",
            },
            {
                "id": "200",
                "active": "1",
                "conceptId": "10",
                "term": "synonyme existant",
                "caseSignificanceId": "cI",
                "acceptabilityId": "ACCEPTABLE",
            },
        ],
        dtype=str,
    )


def test_existing_description_becomes_chg_and_ina(
    tmp_path: Path, descriptions: pd.DataFrame
) -> None:
    source = tmp_path / "delivery.xlsx"
    _delivery(source, [_rep_row()])

    prepared = prepare_delivery_inputs(str(source), descriptions, str(tmp_path))

    chg = pd.read_csv(prepared.batch_files["CHG"], sep=";", dtype=str)
    ina = pd.read_csv(prepared.batch_files["INA"], sep=";", dtype=str)
    rep = pd.read_csv(
        tmp_path / "recreated_inputs" / "Descriptions_Replacements.csv",
        sep=";",
        dtype=str,
    )
    assert rep.empty
    assert chg.loc[0, "Description ID"] == "200"
    assert chg.loc[0, "Case significance"] == "cI"
    assert chg.loc[0, "Acceptability"] == "PREFERRED"
    assert ina.loc[0, "Description ID Or Term"] == "100"

    workbook = load_workbook(prepared.workbook, read_only=True)
    assert workbook["Description Replacements"]["A2"].value is None
    assert workbook["Description Changes"]["A2"].value == "200"
    assert workbook["Description Inactivations"]["A2"].value == "100"


def test_new_term_stays_in_rep(tmp_path: Path, descriptions: pd.DataFrame) -> None:
    source = tmp_path / "delivery.xlsx"
    _delivery(source, [_rep_row(new_id="", new_term="nouveau PT")])

    prepared = prepare_delivery_inputs(str(source), descriptions, str(tmp_path))

    rep = pd.read_csv(prepared.batch_files["REP"], sep=";", dtype=str)
    assert len(rep) == 1
    assert rep.loc[0, "New Translated Term"] == "nouveau PT"
    assert "CHG" not in prepared.batch_files
    assert "INA" not in prepared.batch_files


def test_invalid_row_is_reported_and_valid_rows_are_kept(
    tmp_path: Path, descriptions: pd.DataFrame
) -> None:
    source = tmp_path / "delivery.xlsx"
    invalid_row = _rep_row()
    invalid_row[COL["REP"].index("Concept ID")] = "999"
    _delivery(source, [invalid_row, _rep_row()])

    prepared = prepare_delivery_inputs(str(source), descriptions, str(tmp_path))

    output = tmp_path / "recreated_inputs"
    assert (output / "conversion_report.csv").exists()
    assert (output / "conversion_errors.csv").exists()
    assert (output / "delivery_reconditionne.xlsx").exists()
    assert prepared.rejected_rows == 1
    assert prepared.error_report == str(output / "conversion_errors.csv")

    errors = pd.read_csv(prepared.error_report, sep=";", dtype=str)
    assert len(errors) == 1
    assert errors.loc[0, "conceptId"] == "999"

    chg = pd.read_csv(prepared.batch_files["CHG"], sep=";", dtype=str)
    ina = pd.read_csv(prepared.batch_files["INA"], sep=";", dtype=str)
    assert len(chg) == 1
    assert len(ina) == 1
